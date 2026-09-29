import os
import re
import csv
import openpyxl
import unicodedata
import pythoncom
import win32com.client

def get_sheet_names(file_path):
    if not os.path.exists(file_path):
        return []
    
    if file_path.lower().endswith('.csv'):
        return ['CSV (Arquivo Unico)']
        
    try:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        sheets = wb.sheetnames
        wb.close()
        return sheets
    except Exception:
        return []

def get_sheet_preview_rows(file_path, sheet_name=None, max_rows=15):
    if not os.path.exists(file_path):
        return []
    
    if file_path.lower().endswith('.csv'):
        with open(file_path, 'r', encoding='utf-8', errors='replace') as f:
            sample = f.read(2048)
            delimiter = ';' if ';' in sample else ','
            f.seek(0)
            reader = csv.reader(f, delimiter=delimiter)
            return [row for _, row in zip(range(max_rows), reader)]
    else:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        if sheet_name and sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        else:
            ws = wb.active
        preview = []
        for idx, row in enumerate(ws.iter_rows(values_only=True), 1):
            if idx > max_rows:
                break
            preview.append(list(row))
        wb.close()
        return preview

def get_header_row_suggestions(file_path, sheet_name=None, max_rows=10):
    if not os.path.exists(file_path):
        return [], 0

    rows = []
    if file_path.lower().endswith('.csv'):
        with open(file_path, 'r', encoding='utf-8', errors='replace') as f:
            sample = f.read(2048)
            delimiter = ';' if ';' in sample else ','
            f.seek(0)
            reader = csv.reader(f, delimiter=delimiter)
            rows = [r for _, r in zip(range(max_rows), reader)]
    else:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        if sheet_name and sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        else:
            ws = wb.active
        for idx, row in enumerate(ws.iter_rows(values_only=True), 1):
            if idx > max_rows:
                break
            rows.append(list(row))
        wb.close()

    if not rows:
        return [{'index': 0, 'label': 'Linha 1 (Vazia)'}], 0

    suggestions = []
    best_idx = 0
    best_score = -100

    keywords = ['email', 'e-mail', 'mail', 'nome', 'user', 'usuario', 'usuário', 'id', 'portal', 'código', 'grupo', 'senha', 'password', 'usercode', 'description', 'firstname', 'lastname', 'cpf']

    for idx, row in enumerate(rows):
        non_empty = [str(c).strip() for c in row if c is not None and str(c).strip() != '']
        if not non_empty:
            continue

        count = len(non_empty)
        score = count * 2
        for val in non_empty:
            v_low = val.lower()
            if '\\' in val or '/' in val:
                score -= 15
            else:
                if any(k in v_low for k in keywords):
                    score += 8
                if not v_low.isdigit() and len(val) < 45:
                    score += 2

        sample_items = non_empty[:4]
        sample = ', '.join(sample_items)
        if len(non_empty) > 4:
            sample += f" (+{len(non_empty)-4})"

        label = f"Linha {idx + 1} ({count} colunas): {sample}"
        suggestions.append({'index': idx, 'label': label, 'count': count, 'score': score})

        if score > best_score:
            best_score = score
            best_idx = idx

    if not suggestions:
        suggestions.append({'index': 0, 'label': 'Linha 1 (Padrao)'})

    return suggestions, best_idx

def load_sheet_data(file_path, sheet_name=None, header_row_idx=None):
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Arquivo nao encontrado: {file_path}")

    rows = []
    if file_path.lower().endswith('.csv'):
        with open(file_path, 'r', encoding='utf-8', errors='replace') as f:
            sample = f.read(2048)
            delimiter = ';' if ';' in sample else ','
            f.seek(0)
            reader = csv.reader(f, delimiter=delimiter)
            rows = list(reader)
    else:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        if sheet_name and sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        else:
            ws = wb.active
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()

    while rows and not any(c is not None and str(c).strip() != '' for c in rows[-1]):
        rows.pop()

    if not rows:
        return [], []

    if header_row_idx is None:
        _, best_idx = get_header_row_suggestions(file_path, sheet_name)
        detected_idx = best_idx
    else:
        detected_idx = max(0, min(header_row_idx, len(rows) - 1))

    data_start = detected_idx + 1
    header_row = rows[detected_idx]

    col_has_content = [False] * len(header_row)
    for c_idx in range(len(header_row)):
        if header_row[c_idx] is not None and str(header_row[c_idx]).strip() != '':
            col_has_content[c_idx] = True
        else:
            for d_row in rows[data_start:]:
                if c_idx < len(d_row) and d_row[c_idx] is not None and str(d_row[c_idx]).strip() != '':
                    col_has_content[c_idx] = True
                    break

    columns = []
    valid_col_indices = []
    seen = {}

    for i, c in enumerate(header_row):
        if not col_has_content[i]:
            continue
        
        col_name = str(c).strip() if c is not None and str(c).strip() != '' else f"Coluna_{i+1}"
        if col_name in seen:
            seen[col_name] += 1
            col_name = f"{col_name}_{seen[col_name]}"
        else:
            seen[col_name] = 1

        columns.append(col_name)
        valid_col_indices.append(i)

    records = []
    for row in rows[data_start:]:
        if not any(c is not None and str(c).strip() != '' for c in row):
            continue
            
        item = {}
        for col_name, c_idx in zip(columns, valid_col_indices):
            val = row[c_idx] if c_idx < len(row) else None
            item[col_name] = '' if val is None else str(val).strip()
            
        records.append(item)

    return columns, records

def render_template(template_text, data_dict):
    if not template_text:
        return ""

    rendered = template_text
    for col_name, value in data_dict.items():
        patterns = [
            re.escape(f"#{col_name}#"),
            re.escape(f"#{{{col_name}}}#"),
            re.escape(f"{{{col_name}}}")
        ]
        combined_pattern = '|'.join(patterns)
        rendered = re.sub(combined_pattern, str(value), rendered, flags=re.IGNORECASE)

    return rendered

def text_to_html(plain_text):
    if not plain_text:
        return ""

    formatted = re.sub(r'\*\*(.*?)\*\*', r'<strong>\1</strong>', plain_text)
    url_pattern = r'(https?://[^\s<>\)"]+)'
    formatted = re.sub(url_pattern, r'<a href="\1">\1</a>', formatted)

    linhas = formatted.split('\n')
    paragrafos = []
    bloco_atual = []

    for linha in linhas:
        if linha.strip() == '':
            if bloco_atual:
                paragrafos.append('<p style="margin: 0 0 10px 0;">' + '<br>'.join(bloco_atual) + '</p>')
                bloco_atual = []
        else:
            bloco_atual.append(linha)

    if bloco_atual:
        paragrafos.append('<p style="margin: 0 0 10px 0;">' + '<br>'.join(bloco_atual) + '</p>')

    html_content = ''.join(paragrafos)

    full_html = f"""<html>
<body style="font-family: Calibri, Segoe UI, Arial, sans-serif; font-size: 11pt; color: #1F2937; line-height: 1.5;">
{html_content}
</body>
</html>"""

    return full_html

def normalize_text(text):
    if not text:
        return ""
    text_norm = unicodedata.normalize('NFKD', str(text)).encode('ASCII', 'ignore').decode('ASCII')
    text_clean = re.sub(r'[_\-\.\s]+', ' ', text_norm).strip().lower()
    return text_clean

def find_matching_attachment(folder_path, record_dict, column_name=None, pattern_template=None, file_list_cache=None):
    if not folder_path or not os.path.exists(folder_path):
        return None

    if file_list_cache is not None:
        files = file_list_cache
    else:
        try:
            files = os.listdir(folder_path)
        except Exception:
            return None

    if pattern_template and ('#' in pattern_template or '{' in pattern_template):
        target_raw = render_template(pattern_template, record_dict)
    elif column_name and column_name in record_dict:
        target_raw = str(record_dict.get(column_name, '')).strip()
    else:
        target_raw = ""

    if not target_raw:
        return None

    target_norm = normalize_text(target_raw)
    if not target_norm:
        return None

    for fname in files:
        fpath = os.path.join(folder_path, fname)
        if not os.path.isfile(fpath):
            continue
        stem, _ = os.path.splitext(fname)
        stem_norm = normalize_text(stem)
        if stem_norm == target_norm:
            return fpath

    for fname in files:
        fpath = os.path.join(folder_path, fname)
        if not os.path.isfile(fpath):
            continue
        stem, _ = os.path.splitext(fname)
        stem_norm = normalize_text(stem)
        if stem_norm.startswith(target_norm):
            return fpath

    for fname in files:
        fpath = os.path.join(folder_path, fname)
        if not os.path.isfile(fpath):
            continue
        stem, _ = os.path.splitext(fname)
        stem_norm = normalize_text(stem)
        if target_norm in stem_norm:
            return fpath

    return None

def verify_dynamic_attachments(folder_path, records, column_name=None, pattern_template=None):
    if not folder_path or not os.path.exists(folder_path) or not records:
        return {'total': len(records) if records else 0, 'found': 0, 'missing': len(records) if records else 0, 'missing_list': []}

    try:
        files = os.listdir(folder_path)
    except Exception:
        files = []

    found = 0
    missing = 0
    missing_list = []

    for r in records:
        match = find_matching_attachment(folder_path, r, column_name, pattern_template, file_list_cache=files)
        key_val = r.get(column_name, '') if column_name else str(list(r.values())[0] if r else '')
        if match:
            found += 1
        else:
            missing += 1
            missing_list.append(key_val)

    return {
        'total': len(records),
        'found': found,
        'missing': missing,
        'missing_list': missing_list
    }

def process_emails(records, email_column, subject_template, body_template, 
                   attachments=None, is_draft=True, progress_callback=None, cancel_event=None,
                   group_column=None, group_templates=None,
                   dynamic_folder=None, dynamic_column=None, dynamic_pattern=None,
                   skip_if_missing_dynamic_file=False,
                   cc_template=None, bcc_template=None):
    if not records:
        if progress_callback:
            progress_callback(0, 0, 'error', 'Nenhum registro para processar.', None)
        return {'total': 0, 'success': 0, 'errors': 0, 'skipped': 0}

    attachments = attachments or []
    total = len(records)
    success = 0
    errors = 0
    skipped = 0

    dyn_files_cache = None
    if dynamic_folder and os.path.exists(dynamic_folder):
        try:
            dyn_files_cache = os.listdir(dynamic_folder)
        except Exception:
            dyn_files_cache = []

    pythoncom.CoInitialize()
    try:
        outlook = win32com.client.Dispatch("Outlook.Application")

        for idx, item in enumerate(records, 1):
            if cancel_event and cancel_event.is_set():
                if progress_callback:
                    progress_callback(idx, total, 'warning', 'Processamento cancelado pelo usuario.', None)
                break

            recipient_email = item.get(email_column, '').strip()

            if not recipient_email or '@' not in recipient_email:
                skipped += 1
                if progress_callback:
                    progress_callback(idx, total, 'warning', 
                                      f"Linha {idx}: E-mail invalido ou vazio ({recipient_email}). Registro pulado.", item)
                continue

            cur_subject = subject_template
            cur_body = body_template
            cur_attachments = list(attachments)
            cur_cc = cc_template or ""
            cur_bcc = bcc_template or ""

            if group_column and group_templates:
                group_val = str(item.get(group_column, '')).strip()
                tpl = group_templates.get(group_val)
                if not tpl:
                    for k, v in group_templates.items():
                        if str(k).strip().upper() == group_val.upper():
                            tpl = v
                            break

                if tpl:
                    if tpl.get('subject'):
                        cur_subject = tpl['subject']
                    if tpl.get('body'):
                        cur_body = tpl['body']
                    if 'attachments' in tpl:
                        cur_attachments = list(tpl['attachments'])
                    if 'cc' in tpl:
                        cur_cc = tpl['cc']
                    if 'bcc' in tpl:
                        cur_bcc = tpl['bcc']

            if dynamic_folder and os.path.exists(dynamic_folder):
                dyn_file = find_matching_attachment(
                    dynamic_folder, item, dynamic_column, dynamic_pattern, file_list_cache=dyn_files_cache
                )
                if dyn_file:
                    cur_attachments.append(dyn_file)
                else:
                    if skip_if_missing_dynamic_file:
                        skipped += 1
                        key_info = item.get(dynamic_column or '', '')
                        if progress_callback:
                            progress_callback(idx, total, 'warning', 
                                              f"Linha {idx}: Arquivo individual nao encontrado na pasta para '{key_info}'. Registro pulado.", item)
                        continue

            try:
                subject = render_template(cur_subject, item)
                raw_body = render_template(cur_body, item)
                html_body = text_to_html(raw_body)

                mail = outlook.CreateItem(0)
                mail.To = recipient_email
                mail.Subject = subject
                mail.HTMLBody = html_body

                if cur_cc:
                    rendered_cc = render_template(cur_cc, item).strip()
                    if rendered_cc:
                        mail.CC = rendered_cc

                if cur_bcc:
                    rendered_bcc = render_template(cur_bcc, item).strip()
                    if rendered_bcc:
                        mail.BCC = rendered_bcc

                for att in cur_attachments:
                    if os.path.exists(att):
                        mail.Attachments.Add(att)

                if is_draft:
                    mail.Save()
                    action_name = "Rascunho criado"
                else:
                    mail.Send()
                    action_name = "E-mail enviado"

                success += 1
                if progress_callback:
                    att_names = [os.path.basename(a) for a in cur_attachments]
                    att_str = f" (Anexos: {', '.join(att_names)})" if att_names else ""
                    progress_callback(idx, total, 'success', 
                                      f"[{idx}/{total}] {action_name} para {recipient_email}{att_str}", item)

            except Exception as e:
                errors += 1
                if progress_callback:
                    progress_callback(idx, total, 'error', 
                                      f"[{idx}/{total}] Erro ao processar {recipient_email}: {e}", item)

    finally:
        pythoncom.CoUninitialize()

    return {'total': total, 'success': success, 'errors': errors, 'skipped': skipped}